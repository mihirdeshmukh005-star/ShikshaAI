import csv
import os
import uuid
import json
from pathlib import Path
from pypdf import PdfReader
from .store import add_chunks

SUPPORTED = {".pdf", ".docx", ".pptx", ".txt", ".md", ".csv", ".json"}


def _chunk_text(text, size=1100, overlap=180):
    text = " ".join((text or "").split())
    if not text:
        return []
    chunks = []
    start = 0
    while start < len(text):
        end = min(len(text), start + size)
        if end < len(text):
            boundary = text.rfind(". ", start, end)
            if boundary > start + size // 2:
                end = boundary + 1
        piece = text[start:end].strip()
        if piece:
            chunks.append(piece)
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return chunks


def _extract(path, ext):
    if ext == ".pdf":
        reader = PdfReader(path)
        pages = []
        for page_no, page in enumerate(reader.pages, start=1):
            pages.append((page_no, page.extract_text() or ""))
        return pages
    if ext in {".txt", ".md"}:
        return [(1, Path(path).read_text(encoding="utf-8", errors="ignore"))]
    if ext == ".json":
        raw = Path(path).read_text(encoding="utf-8", errors="ignore")
        try:
            obj = json.loads(raw)
            return [(1, json.dumps(obj, ensure_ascii=False, indent=2))]
        except json.JSONDecodeError:
            return [(1, raw)]
    if ext == ".csv":
        with open(path, newline="", encoding="utf-8", errors="ignore") as f:
            rows = list(csv.reader(f))
        return [(i + 1, " | ".join(row)) for i, row in enumerate(rows)]
    if ext == ".docx":
        from docx import Document
        doc = Document(path)
        return [(1, "\n".join(p.text for p in doc.paragraphs if p.text.strip()))]
    if ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(path)
        pages = []
        for i, slide in enumerate(prs.slides, start=1):
            texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text)
            pages.append((i, "\n".join(texts)))
        return pages
    raise ValueError(f"Unsupported file type: {ext}")


def ingest_material(path, source_name):
    ext = Path(source_name).suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError("Supported materials: PDF, DOCX, PPTX, TXT, MD, CSV and JSON.")
    pages = _extract(path, ext)
    chunks = []
    for page_no, text in pages:
        for idx, piece in enumerate(_chunk_text(text)):
            chunks.append({
                "id": str(uuid.uuid4()),
                "source": source_name,
                "page": page_no,
                "text": piece,
            })
    if not chunks:
        raise ValueError("No readable text was found in the uploaded material.")
    added = add_chunks(chunks)
    return {"pages": len(pages), "chunks": added, "source": source_name}


def ingest_pdf(path, source_name):
    return ingest_material(path, source_name)
